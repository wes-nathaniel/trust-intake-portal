"""
AI Document Explainer — extracts text from an uploaded document (PDF/DOCX/TXT/image)
and asks Claude Sonnet 4.5 (via Emergent LLM Key) to produce a short structured
explanation the client can read directly under the document in the portal.

Public API:
  * extract_text(data: bytes, filename: str, mime: str) -> str
  * generate_explanation(label: str, filename: str, text: str) -> dict
      returns {"what_is", "what_it_does", "why_it_matters", "notarization",
               "client_action", "generated_at"}

Failures never raise — they return a structured error dict so the upload flow
does not break if the LLM is down.
"""
from __future__ import annotations
import io
import json
import logging
import os
import re
from datetime import datetime, timezone

logger = logging.getLogger(__name__)

MAX_TEXT_CHARS = 45000  # Fits comfortably within Claude Sonnet 4.5 context budget.


# --------- text extraction ---------
def _extract_pdf(data: bytes) -> str:
    try:
        import pypdfium2 as pdfium
    except Exception as e:
        logger.warning(f"pypdfium2 unavailable: {e}")
        return ""
    try:
        pdf = pdfium.PdfDocument(data)
        out = []
        for i in range(len(pdf)):
            try:
                page = pdf[i]
                tp = page.get_textpage()
                out.append(tp.get_text_range() or "")
                tp.close()
                page.close()
            except Exception:
                continue
        pdf.close()
        return "\n".join(out).strip()
    except Exception as e:
        logger.warning(f"PDF text extraction failed: {e}")
        return ""


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except Exception as e:
        logger.warning(f"python-docx unavailable: {e}")
        return ""
    try:
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text]
        # Tables
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts).strip()
    except Exception as e:
        logger.warning(f"DOCX text extraction failed: {e}")
        return ""


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc, errors="ignore")
        except Exception:
            continue
    return ""


def extract_text(data: bytes, filename: str, mime: str) -> str:
    fn = (filename or "").lower()
    mt = (mime or "").lower()
    if fn.endswith(".pdf") or "pdf" in mt:
        return _extract_pdf(data)[:MAX_TEXT_CHARS]
    if fn.endswith(".docx") or "wordprocessingml" in mt:
        return _extract_docx(data)[:MAX_TEXT_CHARS]
    if fn.endswith(".txt") or mt.startswith("text/"):
        return _extract_txt(data)[:MAX_TEXT_CHARS]
    # .doc (legacy) and image formats: no reliable text — return empty so the
    # LLM prompt falls back to using only the label + filename.
    return ""


# --------- LLM explainer ---------
SYSTEM_PROMPT = (
    "You are a trust & estate paralegal explaining legal documents to a lay client. "
    "You write short, warm, plain-English explanations — NEVER legal advice. Assume the "
    "client is a U.S. adult unfamiliar with legal jargon. Keep every field concise: 1–3 "
    "sentences each. When you are unsure whether a document requires notarization, say "
    "so honestly (do not invent). Always return VALID JSON matching the schema exactly."
)

USER_TEMPLATE = """The admin has uploaded a document titled "{label}" (filename: {filename}) to a client's trust portal. Explain it to the client.

Return ONLY a JSON object with these exact keys (no markdown, no code fences):
{{
  "what_is": "1-2 sentence definition of the document.",
  "what_it_does": "1-2 sentences on its practical function.",
  "why_it_matters": "1-2 sentences on why the client should care.",
  "notarization": "Say clearly: 'Yes — requires notarization', 'No — no notary needed', or 'Depends — ...' with a short reason. Never guess if unsure.",
  "client_action": "One short line telling the client what (if anything) they must do next. If nothing, say 'No action needed — this is for your records.'"
}}

Document content (may be truncated):
---
{text}
---
"""


def _empty_explanation(error: str = "") -> dict:
    return {
        "what_is": "",
        "what_it_does": "",
        "why_it_matters": "",
        "notarization": "",
        "client_action": "",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "error": error or "",
    }


def _parse_json_lenient(raw: str) -> dict | None:
    if not raw:
        return None
    # Strip common code-fence wrappers Claude sometimes emits despite instructions.
    stripped = raw.strip()
    if stripped.startswith("```"):
        stripped = re.sub(r"^```(?:json)?\s*", "", stripped)
        stripped = re.sub(r"\s*```$", "", stripped)
    try:
        return json.loads(stripped)
    except Exception:
        pass
    # Last-resort: grab first {...} block
    m = re.search(r"\{.*\}", stripped, re.DOTALL)
    if m:
        try:
            return json.loads(m.group(0))
        except Exception:
            return None
    return None


async def generate_explanation(label: str, filename: str, text: str,
                                case_context: dict | None = None) -> dict:
    """Call Claude Sonnet 4.5 and return a structured explanation dict.
    Never raises — logs and returns an error-tagged dict on failure.

    When `case_context` is supplied (grantors, trustees, successors, beneficiaries,
    assets, trust_type, client_name), we emit a **deeply personalized** trust
    explanation using client + fiduciary names, share percentages, and asset
    descriptions in plain English. Falls back to the generic 5-field schema when
    no context is passed."""
    api_key = os.environ.get("EMERGENT_LLM_KEY", "").strip()
    if not api_key:
        return _empty_explanation("EMERGENT_LLM_KEY not configured")

    text_for_prompt = (text or "").strip()
    if not text_for_prompt:
        text_for_prompt = (
            "(No machine-readable text could be extracted from this file. Base your "
            "explanation on the document label and filename only, and clearly note that "
            "the actual contents were not analyzed.)"
        )

    personalized = bool(case_context)
    if personalized:
        user_prompt = _personalized_template(label, filename, text_for_prompt, case_context)
    else:
        user_prompt = USER_TEMPLATE.format(
            label=(label or "(unlabeled)").strip(),
            filename=filename or "unknown",
            text=text_for_prompt,
        )

    try:
        from emergentintegrations.llm.chat import LlmChat, UserMessage
        chat = LlmChat(
            api_key=api_key,
            session_id=f"doc-explainer-{datetime.now(timezone.utc).timestamp()}",
            system_message=(PERSONALIZED_SYSTEM_PROMPT if personalized else SYSTEM_PROMPT),
        ).with_model("anthropic", "claude-sonnet-4-5-20250929")
        raw = await chat.send_message(UserMessage(text=user_prompt))
    except Exception as e:
        logger.error(f"AI explainer call failed: {e}")
        return _empty_explanation(f"LLM call failed: {e}")

    raw_str = raw if isinstance(raw, str) else str(raw)
    parsed = _parse_json_lenient(raw_str)
    if not parsed:
        logger.error(f"AI explainer returned unparsable output: {raw_str[:400]}")
        return _empty_explanation("Model returned non-JSON output")

    if personalized:
        return {
            "personalized":       True,
            "greeting":           str(parsed.get("greeting") or "").strip(),
            "grantor_intro":      str(parsed.get("grantor_intro") or "").strip(),
            "trustee_roles":      str(parsed.get("trustee_roles") or "").strip(),
            "beneficiaries":      str(parsed.get("beneficiaries") or "").strip(),
            "assets_plain":       str(parsed.get("assets_plain") or "").strip(),
            "trust_type_meaning": str(parsed.get("trust_type_meaning") or "").strip(),
            "next_steps":         str(parsed.get("next_steps") or "").strip(),
            # Keep the legacy 5 fields present so downstream consumers that read
            # them (admin card, existing UI chips) don't blank out.
            "what_is":            str(parsed.get("what_is") or "").strip(),
            "what_it_does":       str(parsed.get("what_it_does") or "").strip(),
            "why_it_matters":     str(parsed.get("why_it_matters") or "").strip(),
            "notarization":       str(parsed.get("notarization") or "").strip(),
            "client_action":      str(parsed.get("client_action") or parsed.get("next_steps") or "").strip(),
            "generated_at":       datetime.now(timezone.utc).isoformat(),
            "error":              "",
        }

    return {
        "what_is":       str(parsed.get("what_is") or "").strip(),
        "what_it_does":  str(parsed.get("what_it_does") or "").strip(),
        "why_it_matters": str(parsed.get("why_it_matters") or "").strip(),
        "notarization":  str(parsed.get("notarization") or "").strip(),
        "client_action": str(parsed.get("client_action") or "").strip(),
        "generated_at":  datetime.now(timezone.utc).isoformat(),
        "error":         "",
    }


# --------- Personalized trust explanation (deep context) ---------
PERSONALIZED_SYSTEM_PROMPT = (
    "You are writing a warm, deeply personalized plain-English explanation of a "
    "client's own trust document. You are NOT giving legal advice — you are simply "
    "helping this specific person understand what THEIR trust says. Use the client's "
    "first name. Refer to each grantor, trustee, successor trustee, and beneficiary "
    "by name. State each beneficiary's exact share when known. Describe assets in "
    "plain, human language (e.g., \"the family home at 123 Maple Street\", not "
    "\"the residential real property identified in Schedule A\"). NEVER use words "
    "like 'settlor', 'corpus', 'per stirpes', 'fiduciary', 'in terrorem', "
    "'spendthrift', 'hereinafter', 'aforementioned' — if a legal term is essential, "
    "define it inline in parentheses using everyday words. Write like you are "
    "talking to a family member who has never seen a trust before. Warm, human, "
    "confident. Never use corporate boilerplate. Always return VALID JSON only."
)


def _fmt_people(rows, roles_key: str = "role") -> str:
    """Format grantors/trustees/beneficiaries lists into a readable line each."""
    if not rows:
        return "(none listed)"
    out = []
    for r in rows:
        if not isinstance(r, dict):
            continue
        name = " ".join(x for x in [r.get("first_name"), r.get("middle_name"), r.get("last_name")] if x).strip()
        if not name:
            name = r.get("name") or r.get("full_name") or "(unnamed)"
        parts = [name]
        rel = r.get("relationship")
        if rel:
            parts.append(f"({rel})")
        share = r.get("share_percent") or r.get("share")
        if share not in (None, ""):
            try:
                pct = float(share)
                parts.append(f"— {pct:g}% share")
            except Exception:
                parts.append(f"— {share} share")
        role = r.get(roles_key)
        if role:
            parts.append(f"[{role}]")
        out.append(" ".join(parts))
    return "\n".join(f"  • {line}" for line in out)


def _fmt_assets(assets) -> str:
    if not assets:
        return "(none listed on file)"
    out = []
    for a in assets:
        if not isinstance(a, dict):
            continue
        label = a.get("description") or a.get("name") or a.get("type") or "(asset)"
        val = a.get("estimated_value") or a.get("value")
        acct = a.get("account_number") or a.get("account") or ""
        loc = a.get("location") or a.get("address") or ""
        extra_bits = []
        if val not in (None, ""):
            extra_bits.append(f"est. value ${val}")
        if loc:
            extra_bits.append(f"at {loc}")
        if acct:
            extra_bits.append(f"acct …{str(acct)[-4:]}")
        extra = f" ({'; '.join(extra_bits)})" if extra_bits else ""
        out.append(f"  • {label}{extra}")
    return "\n".join(out) or "(none listed on file)"


def _personalized_template(label: str, filename: str, text: str, ctx: dict) -> str:
    client_name = (ctx.get("client_name") or "").strip()
    first_name = client_name.split(" ")[0] if client_name else "there"
    trust_type = (ctx.get("trust_type") or "").strip() or "revocable living trust"
    trust_name = (ctx.get("trust_name") or "").strip()

    grantors_block = _fmt_people(ctx.get("grantors") or [])
    trustees_block = _fmt_people(ctx.get("trustees") or [])
    successors_block = _fmt_people(ctx.get("successor_trustees") or [])
    beneficiaries_block = _fmt_people(ctx.get("beneficiaries") or [])
    assets_block = _fmt_assets(ctx.get("assets") or [])

    return f"""You are writing the plain-English side of {first_name}'s trust preview. This document is titled "{label}" (filename: {filename}).

CLIENT & TRUST CONTEXT (use every piece of this — names, percentages, addresses. Do NOT invent details not listed here):
- Client's full name: {client_name or '(unknown)'}
- Trust name: {trust_name or '(same as client name in most cases)'}
- Trust type: {trust_type}

Grantors (the people creating the trust):
{grantors_block}

Trustees (managing the trust day-to-day):
{trustees_block}

Successor Trustees (steps in on incapacity or death):
{successors_block}

Beneficiaries (who receives the assets, and how much):
{beneficiaries_block}

Assets currently listed in the intake:
{assets_block}

DOCUMENT TEXT (may be truncated — use it for direct quotes but rely on the context above for names/percentages):
---
{text}
---

Return ONLY this JSON (no markdown, no code fences). Every field must be plain-English HTML (<p>, <ul>, <li>, <strong>). NO legal jargon. Address {first_name} directly by name where it feels natural:

{{
  "greeting":           "<p>Hi {first_name} — here's your trust in plain English. Take your time reading through this and jot down any questions.</p>",
  "grantor_intro":      "<p>As the <strong>Grantor</strong> (the person creating this trust), YOU are... </p><ul><li>...powers you keep while alive and well</li><li>...what happens if you become unable to manage things</li><li>...what happens at your death</li></ul>",
  "trustee_roles":      "<p>Your <strong>Trustee</strong> right now is [NAME] — that's the person managing the trust. Your <strong>Successor Trustee</strong>, [NAME], steps in when...</p><p>Their job is to...</p>",
  "beneficiaries":      "<p>These are the people receiving from your trust:</p><ul><li><strong>[Name]</strong> — [relationship] — [share%] of...</li>...</ul><p>Optional 1-line note on how the shares work together (per stirpes / equal / etc — but explained in plain words, not with the Latin term).</p>",
  "assets_plain":       "<p>Here's what's going into your trust so far, in everyday language:</p><ul><li>[plain description of each asset — e.g., 'your home at 123 Maple St, worth about $X']</li>...</ul><p>Anything you own that's NOT on this list won't be inside the trust unless you retitle it. Talk to us before making that change.</p>",
  "trust_type_meaning": "<p>Because this is a <strong>{trust_type}</strong>, it means for you specifically: [what it does + doesn't do, in 2-4 short bullet points tailored to this client's situation — not a generic definition].</p>",
  "next_steps":         "<p>Here's what happens now, {first_name}:</p><ol><li>Read through your draft on the left, pane by pane.</li><li>[personalized next step based on where they are — e.g., add missing beneficiary detail, sign notarization, etc.]</li><li>[one more specific action]</li></ol><p>When you're ready, use the buttons at the bottom to schedule a call or complete your review.</p>",

  "what_is":            "1 short line — same voice as above.",
  "what_it_does":       "1 short line.",
  "why_it_matters":     "1 short line about how this specific trust helps {first_name}.",
  "notarization":       "Say clearly whether THIS trust requires notarization ('Yes — requires notarization by...' / 'No' / 'Depends — ...').",
  "client_action":      "One short line telling {first_name} exactly what to do next."
}}
"""
